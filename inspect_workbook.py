import json
import os
import re
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime, time

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


SOURCE = r"C:\Users\lenovo\Downloads\Expesne Tracker sheet.xlsx"
OUTPUT = r"C:\Users\lenovo\Documents\ChatGPT\expense management\workbook_audit.json"


def serialise(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bytes):
        return f"<{len(value)} bytes>"
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def color_value(color):
    if color is None:
        return None
    return {
        "type": getattr(color, "type", None),
        "rgb": getattr(color, "rgb", None),
        "indexed": getattr(color, "indexed", None),
        "theme": getattr(color, "theme", None),
        "tint": getattr(color, "tint", None),
    }


def chart_summary(chart):
    title = None
    try:
        title = chart.title.tx.rich.p[0].r[0].t
    except Exception:
        try:
            title = chart.title.tx.rich.p[0].endParaRPr.t
        except Exception:
            title = str(chart.title) if chart.title else None
    def source_ref(obj):
        if obj is None:
            return None
        for child_name in ("numRef", "strRef"):
            child = getattr(obj, child_name, None)
            if child is not None and getattr(child, "f", None):
                return child.f
        return None

    series = []
    for ser in getattr(chart, "ser", []) or []:
        entry = {}
        for key in ("val", "cat", "xVal", "yVal"):
            obj = getattr(ser, key, None)
            ref = source_ref(obj)
            if ref:
                entry[key] = ref
        tx = getattr(ser, "tx", None)
        if tx is not None:
            if getattr(tx, "strRef", None) is not None:
                entry["tx"] = tx.strRef.f
            elif getattr(tx, "v", None) is not None:
                entry["tx"] = tx.v
        series.append(entry)
    anchor = getattr(chart, "anchor", None)
    anchor_data = None
    if anchor is not None:
        anchor_data = {
            "from": {
                "col": getattr(getattr(anchor, "_from", None), "col", None),
                "row": getattr(getattr(anchor, "_from", None), "row", None),
            },
            "to": {
                "col": getattr(getattr(anchor, "to", None), "col", None),
                "row": getattr(getattr(anchor, "to", None), "row", None),
            },
        }
    return {
        "type": chart.__class__.__name__,
        "title": title,
        "style": getattr(chart, "style", None),
        "height": getattr(chart, "height", None),
        "width": getattr(chart, "width", None),
        "anchor": anchor_data,
        "series": series,
        "legend_position": getattr(getattr(chart, "legend", None), "position", None),
        "data_labels": serialise(getattr(chart, "dLbls", None)),
    }


def drawing_anchor_summary(anchor):
    if isinstance(anchor, str):
        return anchor
    result = {"type": anchor.__class__.__name__}
    for label, attr in (("from", "_from"), ("to", "to")):
        marker = getattr(anchor, attr, None)
        if marker is not None:
            result[label] = {
                "col": getattr(marker, "col", None),
                "colOff": getattr(marker, "colOff", None),
                "row": getattr(marker, "row", None),
                "rowOff": getattr(marker, "rowOff", None),
            }
    return result


wb_formula = load_workbook(SOURCE, data_only=False, read_only=False)
wb_values = load_workbook(SOURCE, data_only=True, read_only=False)

report = {
    "source": SOURCE,
    "file_size": os.path.getsize(SOURCE),
    "workbook": {
        "sheetnames": wb_formula.sheetnames,
        "active_sheet": wb_formula.active.title,
        "calculation": serialise(getattr(wb_formula, "calculation", None)),
        "properties": {k: serialise(v) for k, v in vars(wb_formula.properties).items() if not k.startswith("_")},
        "security": serialise(wb_formula.security),
        "defined_names": [],
    },
    "sheets": [],
    "package": {},
}

for name, defn in wb_formula.defined_names.items():
    report["workbook"]["defined_names"].append({
        "name": name,
        "attr_text": defn.attr_text,
        "hidden": defn.hidden,
        "localSheetId": defn.localSheetId,
    })

for ws in wb_formula.worksheets:
    wsv = wb_values[ws.title]
    used_cells = []
    formulas = []
    strings = []
    numbers = []
    errors = []
    style_counts = Counter()
    number_formats = Counter()
    fill_colors = Counter()
    font_colors = Counter()
    alignments = Counter()
    borders = Counter()
    row_nonempty = defaultdict(int)
    col_nonempty = defaultdict(int)
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            val = serialise(cell.value)
            cached = serialise(wsv[cell.coordinate].value)
            item = {
                "cell": cell.coordinate,
                "value": val,
                "cached_value": cached,
                "data_type": cell.data_type,
                "number_format": cell.number_format,
                "style_id": cell.style_id,
            }
            if cell.hyperlink:
                item["hyperlink"] = cell.hyperlink.target or cell.hyperlink.location
            if cell.comment:
                item["comment"] = {"author": cell.comment.author, "text": cell.comment.text}
            used_cells.append(item)
            row_nonempty[cell.row] += 1
            col_nonempty[cell.column] += 1
            style_counts[cell.style_id] += 1
            number_formats[cell.number_format] += 1
            fill_colors[str(color_value(cell.fill.fgColor))] += 1
            font_colors[str(color_value(cell.font.color))] += 1
            alignments[str((cell.alignment.horizontal, cell.alignment.vertical, cell.alignment.wrap_text, cell.alignment.text_rotation))] += 1
            borders[str((cell.border.left.style, cell.border.right.style, cell.border.top.style, cell.border.bottom.style))] += 1
            if cell.data_type == "f":
                formulas.append(item)
            elif cell.data_type == "e":
                errors.append(item)
            elif isinstance(cell.value, str):
                strings.append(item)
            elif isinstance(cell.value, (int, float)):
                numbers.append(item)

    validations = []
    for dv in ws.data_validations.dataValidation:
        validations.append({
            "type": dv.type,
            "sqref": str(dv.sqref),
            "formula1": dv.formula1,
            "formula2": dv.formula2,
            "allow_blank": dv.allowBlank,
            "show_error": dv.showErrorMessage,
            "show_input": dv.showInputMessage,
            "promptTitle": dv.promptTitle,
            "prompt": dv.prompt,
            "errorTitle": dv.errorTitle,
            "error": dv.error,
        })

    conditional = []
    for ranges, rules in ws.conditional_formatting._cf_rules.items():
        for rule in rules:
            conditional.append({
                "sqref": str(ranges.sqref),
                "type": rule.type,
                "operator": rule.operator,
                "formula": list(rule.formula or []),
                "priority": rule.priority,
                "stopIfTrue": rule.stopIfTrue,
                "dxfId": rule.dxfId,
            })

    tables = []
    for table in ws.tables.values():
        tables.append({
            "name": table.name,
            "displayName": table.displayName,
            "ref": table.ref,
            "style": serialise(table.tableStyleInfo),
            "autoFilter": serialise(table.autoFilter),
            "columns": [{"id": c.id, "name": c.name, "totalsRowFunction": c.totalsRowFunction} for c in table.tableColumns],
        })

    merged = [str(rng) for rng in ws.merged_cells.ranges]
    hidden_rows = [i for i, dim in ws.row_dimensions.items() if dim.hidden]
    hidden_cols = [k for k, dim in ws.column_dimensions.items() if dim.hidden]
    outlined_rows = {str(i): dim.outlineLevel for i, dim in ws.row_dimensions.items() if dim.outlineLevel}
    outlined_cols = {k: dim.outlineLevel for k, dim in ws.column_dimensions.items() if dim.outlineLevel}
    row_heights = {str(i): dim.height for i, dim in ws.row_dimensions.items() if dim.height is not None}
    col_widths = {k: dim.width for k, dim in ws.column_dimensions.items() if dim.width is not None}

    report["sheets"].append({
        "title": ws.title,
        "state": ws.sheet_state,
        "dimensions": ws.calculate_dimension(),
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "sheet_properties": serialise(ws.sheet_properties),
        "sheet_view": {
            "showGridLines": ws.sheet_view.showGridLines,
            "zoomScale": ws.sheet_view.zoomScale,
            "rightToLeft": ws.sheet_view.rightToLeft,
            "freeze_panes": serialise(ws.freeze_panes),
            "selected_cell": serialise(getattr(ws.sheet_view.selection[0], "activeCell", None)) if ws.sheet_view.selection else None,
        },
        "print": {
            "print_area": serialise(ws.print_area),
            "print_title_rows": ws.print_title_rows,
            "print_title_cols": ws.print_title_cols,
            "page_setup": serialise(ws.page_setup),
            "page_margins": serialise(ws.page_margins),
        },
        "auto_filter": str(ws.auto_filter.ref) if ws.auto_filter.ref else None,
        "merged_cells": merged,
        "hidden_rows": hidden_rows,
        "hidden_columns": hidden_cols,
        "outlined_rows": outlined_rows,
        "outlined_columns": outlined_cols,
        "row_heights": row_heights,
        "column_widths": col_widths,
        "validations": validations,
        "conditional_formatting": conditional,
        "tables": tables,
        "charts": [chart_summary(c) for c in ws._charts],
        "images": [{"width": i.width, "height": i.height, "format": getattr(i, "format", None), "anchor": drawing_anchor_summary(i.anchor)} for i in ws._images],
        "protection": serialise(ws.protection),
        "style_counts": dict(style_counts),
        "number_formats": dict(number_formats),
        "fill_colors": dict(fill_colors),
        "font_colors": dict(font_colors),
        "alignments": dict(alignments),
        "borders": dict(borders),
        "nonempty_rows": dict(row_nonempty),
        "nonempty_columns": {get_column_letter(k): v for k, v in col_nonempty.items()},
        "formula_count": len(formulas),
        "formulas": formulas,
        "error_cells": errors,
        "string_cells": strings,
        "number_cells": numbers,
        "all_used_cells": used_cells,
    })

with zipfile.ZipFile(SOURCE) as zf:
    names = zf.namelist()
    report["package"] = {
        "file_count": len(names),
        "files": names,
        "has_vba": any("vbaProject" in n for n in names),
        "pivot_files": [n for n in names if "pivot" in n.lower()],
        "chart_files": [n for n in names if n.startswith("xl/charts/")],
        "drawing_files": [n for n in names if n.startswith("xl/drawings/")],
        "external_links": [n for n in names if n.startswith("xl/externalLinks/")],
        "slicer_files": [n for n in names if "slicer" in n.lower()],
        "custom_xml": [n for n in names if n.startswith("customXml/")],
    }

with open(OUTPUT, "w", encoding="utf-8") as f:
    json.dump(report, f, indent=2, ensure_ascii=False, default=serialise)

print(json.dumps({
    "output": OUTPUT,
    "sheets": [{"title": s["title"], "state": s["state"], "dimensions": s["dimensions"], "cells": len(s["all_used_cells"]), "formulas": s["formula_count"], "charts": len(s["charts"]), "validations": len(s["validations"]), "conditional_rules": len(s["conditional_formatting"])} for s in report["sheets"]],
    "defined_names": len(report["workbook"]["defined_names"]),
    "package": report["package"],
}, indent=2))
